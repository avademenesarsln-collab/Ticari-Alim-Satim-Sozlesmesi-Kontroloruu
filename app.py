import streamlit as st
import google.generativeai as genai
import PyPDF2
import docx
import io

# Sayfa Ayarları
st.set_page_config(page_title="Ticari Sözleşme Asistanı", page_icon="⚖️", layout="wide")

st.title("⚖️ Ticari Mal Alım Satım Sözleşmesi Asistanı")
st.write("Sözleşme taslağınızı PDF veya Word olarak yükleyin, yapay zeka hukuki riskleri bir denetim tablosu olarak analiz etsin.")

# Hafıza (Session State) Ayarı: İndir butonuna basınca analizin ekrandan silinmesini engeller
if "analiz_raporu" not in st.session_state:
    st.session_state.analiz_raporu = None

# Şifre Çekme ve Model Ayarlama
try:
    GEMINI_API_KEY = st.secrets["GEMINI_API_KEY"]
    genai.configure(api_key=GEMINI_API_KEY)
    model = genai.GenerativeModel('gemini-1.5-pro-latest')
except Exception as e:
    st.error("API Anahtarı bulunamadı. Lütfen Streamlit Secrets ayarlarınızı kontrol edin.")

# 1. DOSYA YÜKLEME ALANI
yuklenen_dosya = st.file_uploader("Sözleşme Dosyasını Yükleyin (PDF veya DOCX formatında)", type=["pdf", "docx"])

sozlesme_metni = ""

# 2. DOSYA OKUMA İŞLEMİ
if yuklenen_dosya is not None:
    if yuklenen_dosya.name.endswith('.pdf'):
        try:
            pdf_okuyucu = PyPDF2.PdfReader(yuklenen_dosya)
            for sayfa in pdf_okuyucu.pages:
                if sayfa.extract_text():
                    sozlesme_metni += sayfa.extract_text() + "\n"
            st.success("✅ PDF başarıyla okundu! Metni aşağıda inceleyebilir veya düzenleyebilirsiniz.")
        except Exception as e:
            st.error(f"PDF okuma hatası: {e}")
            
    elif yuklenen_dosya.name.endswith('.docx'):
        try:
            doc = docx.Document(yuklenen_dosya)
            for paragraf in doc.paragraphs:
                sozlesme_metni += paragraf.text + "\n"
            st.success("✅ Word dosyası başarıyla okundu! Metni aşağıda inceleyebilir veya düzenleyebilirsiniz.")
        except Exception as e:
            st.error(f"Word okuma hatası: {e}")

# 3. METİN KUTUSU
guncel_metin = st.text_area("Sözleşme Metni (İsterseniz düzenleyebilir veya doğrudan buraya yapıştırabilirsiniz):", value=sozlesme_metni, height=300)

# 4. ANALİZ BUTONU
if st.button("Sözleşmeyi Hukuken Analiz Et"):
    if guncel_metin.strip():
        with st.spinner("Sözleşme riskleri hesaplanıyor ve denetim tablosu oluşturuluyor..."):
            
            prompt = f"""
            Sen Türkiye'de görev yapan, İstanbul Barosuna kayıtlı uzman bir Ticaret Hukuku Avukatı ve İç Denetim/Risk Yönetimi uzmanısın.
            Aşağıdaki ticari mal alım satım sözleşmesi metnini incele. 
            
            Şu 4 ana başlıkta inceleme yap:
            1. SÜRELER VE İHBARLAR: TTK m.23 ayıp ihbar süreleri ve TBK zamanaşımı süreleri.
            2. YETKİLİ MAHKEME VE ÇÖZÜM: HMK m.17 tacirler arası yetki sözleşmesi geçerlilik şartları.
            3. ALACAK FAİZİ: TTK m.1530 ticari işlerde temerrüt faizi oranları.
            4. EDİMLER VE HASARIN GEÇİŞİ: Teslim şartları ve TBK hasar sorumluluğu.

            LÜTFEN ÇIKTIYI SADECE AŞAĞIDAKİ GİBİ BİR MARKDOWN TABLOSU FORMATINDA VER. Uzun paragraflar yazma.
            
            | İnceleme Konusu | Sözleşmedeki Mevcut Durum | Hukuki Risk Seviyesi (Düşük/Orta/Yüksek) | İlgili Mevzuat (TTK/TBK/HMK) | Revizyon Önerisi ve Çözüm |
            | :--- | :--- | :--- | :--- | :--- |
            | (Konu) | (Sözleşmede ne yazıyor) | (Risk derecesi) | (Kanun maddesi) | (Nasıl düzeltilmeli) |

            Tablonun altına, sözleşmenin genel risk durumunu özetleyen en fazla 3 cümlelik kısa bir "Yönetici Özeti (Executive Summary)" ekle.

            İncelenecek Sözleşme Taslağı:
            {guncel_metin}
            """

            try:
                response = model.generate_content(
                    prompt,
                    generation_config=genai.types.GenerationConfig(
                        temperature=0.1,
                        max_output_tokens=8192,
                    )
                )
                # Analiz sonucunu hafızaya kaydediyoruz
                st.session_state.analiz_raporu = response.text
                
            except Exception as e:
                st.error(f"Yapay Zeka API Hatası: {e}")
                
    else:
        st.warning("Lütfen analiz edilecek bir dosya yükleyin veya kutuya sözleşme metnini girin.")

# 5. SONUÇ EKRANI VE İNDİRME BUTONU
if st.session_state.analiz_raporu:
    st.success("Hukuki Analiz Tamamlandı!")
    st.markdown(st.session_state.analiz_raporu)
    
    # Arka planda Word dosyası oluşturma işlemi
    doc = docx.Document()
    doc.add_heading('Sözleşme Hukuki Risk Denetim Raporu', 0)
    doc.add_paragraph("Bu rapor, Yapay Zeka Destekli Ticari Sözleşme Asistanı tarafından oluşturulmuştur.\n")
    doc.add_paragraph(st.session_state.analiz_raporu)
    
    # Dosyayı bilgisayara indirmek için sanal hafızada (BytesIO) tutuyoruz
    bio = io.BytesIO()
    doc.save(bio)
    
    st.download_button(
        label="📄 Raporu Word (.docx) Olarak İndir",
        data=bio.getvalue(),
        file_name="sozlesme_risk_raporu.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
